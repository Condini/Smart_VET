import copy
import hashlib
import json
import re
import threading
import unicodedata
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


PATIENT_REQUIRED_FIELDS = {
    "id",
    "nome",
    "esp",
    "raca",
    "idade",
    "peso",
    "temp",
    "fc",
    "evolucao",
    "historico",
    "sintomas",
    "semiologia_lesoes",
    "hist_ectoparasiticidas",
    "protocolo",
    "versao_prot",
    "laudo",
    "limitacoes_exame",
    "diferenciais",
    "gravidade",
    "prescricao",
    "validacao_dose",
    "manejo_ambiental",
    "reavaliacao",
}

MEDICATION_REQUIRED_COLUMNS = {
    "ID",
    "Nome Comercial",
    "Princípio Ativo",
    "Classe",
    "Espécies",
    "Dose Min (mg/kg)",
    "Dose Max (mg/kg)",
    "Via",
    "Intervalo",
    "Contraindicações",
    "Alerta FDA",
}

MEDICATION_ALIAS_STOPWORDS = {
    "acido",
    "benzoato",
    "bissulfato",
    "citrato",
    "cloridrato",
    "fosfato",
    "hidratado",
    "hiclato",
    "imunoglobulina",
    "maleato",
    "pomada",
    "solucao",
    "sodica",
    "sodico",
    "sulfato",
    "tartarato",
}


class ClinicalDataError(RuntimeError):
    pass


class PatientNotFoundError(LookupError):
    pass


def normalize_search_text(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(character for character in text if not unicodedata.combining(character))
    return " ".join(re.sub(r"[^a-zA-Z0-9]+", " ", text).lower().split())


class ClinicalRepository:
    def __init__(self, data_dir, reload_on_change=True):
        self.data_dir = Path(data_dir)
        self.patients_path = self.data_dir / "pacientes_clinicos.json"
        self.medications_path = self.data_dir / "medicamentos.xlsx"
        self.reload_on_change = reload_on_change
        self._lock = threading.RLock()
        self._signature = None
        self._patients = {}
        self._medications = []
        self._medication_aliases = []
        self._data_version = "unloaded"
        self._loaded_at = None

    def load(self, force=False):
        with self._lock:
            signature = self._source_signature()
            if not force and self._signature == signature:
                return

            patients = self._load_patients()
            medications = self._load_medications()

            self._patients = patients
            self._medications = medications
            self._medication_aliases = [
                (medication, self._aliases_for_medication(medication))
                for medication in medications
            ]
            self._signature = signature
            self._data_version = self._calculate_data_version()
            self._loaded_at = datetime.now(timezone.utc).isoformat()

    def list_patients(self):
        self._ensure_loaded()
        with self._lock:
            return [
                {
                    "id": patient["id"],
                    "name": patient["nome"],
                    "species": patient["esp"],
                    "breed": patient["raca"],
                }
                for patient in sorted(self._patients.values(), key=lambda item: item["id"])
            ]

    def get_context(self, patient_id):
        self._ensure_loaded()
        normalized_id = str(patient_id or "").strip().upper()

        with self._lock:
            patient = self._patients.get(normalized_id)
            if patient is None:
                raise PatientNotFoundError(f"Patient '{normalized_id}' was not found")

            medications = self._match_medications(patient.get("prescricao", ""))
            return {
                "patient": copy.deepcopy(patient),
                "medications": copy.deepcopy(medications),
                "metadata": {
                    "patient_id": patient["id"],
                    "patient_name": patient["nome"],
                    "protocol_version": patient["versao_prot"],
                    "data_version": self._data_version,
                    "matched_medication_ids": [item["id"] for item in medications],
                },
            }

    def metadata(self):
        self._ensure_loaded()
        with self._lock:
            return {
                "patient_count": len(self._patients),
                "medication_count": len(self._medications),
                "data_version": self._data_version,
                "loaded_at": self._loaded_at,
            }

    def _ensure_loaded(self):
        if self._signature is None:
            self.load(force=True)
        elif self.reload_on_change and self._signature != self._source_signature():
            self.load(force=True)

    def _source_signature(self):
        signature = []
        for path in (self.patients_path, self.medications_path):
            if not path.is_file():
                raise ClinicalDataError(f"Clinical data file was not found: {path.name}")
            stat = path.stat()
            signature.append((path.name, stat.st_mtime_ns, stat.st_size))
        return tuple(signature)

    def _load_patients(self):
        try:
            with self.patients_path.open("r", encoding="utf-8") as stream:
                records = json.load(stream)
        except (OSError, json.JSONDecodeError) as exc:
            raise ClinicalDataError("Could not load pacientes_clinicos.json") from exc

        if not isinstance(records, list) or not records:
            raise ClinicalDataError("Patient data must be a non-empty JSON list")

        patients = {}
        for index, patient in enumerate(records, start=1):
            if not isinstance(patient, dict):
                raise ClinicalDataError(f"Patient record {index} must be an object")

            missing_fields = PATIENT_REQUIRED_FIELDS - set(patient)
            if missing_fields:
                raise ClinicalDataError(
                    f"Patient record {index} is missing fields: {sorted(missing_fields)}"
                )

            patient_id = str(patient["id"]).strip().upper()
            if not re.fullmatch(r"VET-\d{3}", patient_id):
                raise ClinicalDataError(f"Invalid patient ID: {patient_id}")
            if patient_id in patients:
                raise ClinicalDataError(f"Duplicate patient ID: {patient_id}")
            if not isinstance(patient["diferenciais"], list):
                raise ClinicalDataError(f"Patient {patient_id} has invalid differential diagnoses")

            normalized_patient = dict(patient)
            normalized_patient["id"] = patient_id
            patients[patient_id] = normalized_patient

        return patients

    def _load_medications(self):
        try:
            workbook = load_workbook(self.medications_path, read_only=True, data_only=True)
            worksheet = workbook.active
            rows = list(worksheet.iter_rows(values_only=True))
            workbook.close()
        except (OSError, ValueError) as exc:
            raise ClinicalDataError("Could not load medicamentos.xlsx") from exc

        if len(rows) < 2:
            raise ClinicalDataError("Medication workbook must contain a header and data rows")

        headers = [str(value).strip() if value is not None else "" for value in rows[0]]
        missing_columns = MEDICATION_REQUIRED_COLUMNS - set(headers)
        if missing_columns:
            raise ClinicalDataError(
                f"Medication workbook is missing columns: {sorted(missing_columns)}"
            )

        medications = []
        medication_ids = set()
        for row_number, values in enumerate(rows[1:], start=2):
            medication = dict(zip(headers, values))
            medication_id = str(medication["ID"] or "").strip().upper()
            if not medication_id:
                raise ClinicalDataError(f"Medication row {row_number} has no ID")
            if medication_id in medication_ids:
                raise ClinicalDataError(f"Duplicate medication ID: {medication_id}")

            dose_min = float(medication["Dose Min (mg/kg)"])
            dose_max = float(medication["Dose Max (mg/kg)"])
            if dose_min > dose_max:
                raise ClinicalDataError(f"Medication {medication_id} has an invalid dose range")

            medication_ids.add(medication_id)
            medications.append(
                {
                    "id": medication_id,
                    "commercial_name": medication["Nome Comercial"],
                    "active_ingredient": medication["Princípio Ativo"],
                    "drug_class": medication["Classe"],
                    "species": medication["Espécies"],
                    "dose_min_mg_kg": dose_min,
                    "dose_max_mg_kg": dose_max,
                    "route": medication["Via"],
                    "interval": medication["Intervalo"],
                    "contraindications": medication["Contraindicações"],
                    "safety_alert": medication["Alerta FDA"],
                }
            )

        return medications

    def _aliases_for_medication(self, medication):
        aliases = set()
        for value in (medication["commercial_name"], medication["active_ingredient"]):
            normalized_value = normalize_search_text(value)
            if normalized_value:
                aliases.add(normalized_value)

            for component in re.split(r"[/+()]", str(value or "")):
                normalized_component = normalize_search_text(component)
                if len(normalized_component) >= 4:
                    aliases.add(normalized_component)

            for token in normalized_value.split():
                if len(token) >= 5 and token not in MEDICATION_ALIAS_STOPWORDS:
                    aliases.add(token)

        return sorted(aliases, key=len, reverse=True)

    def _match_medications(self, prescription):
        normalized_prescription = f" {normalize_search_text(prescription)} "
        matches = []

        for medication, aliases in self._medication_aliases:
            if any(f" {alias} " in normalized_prescription for alias in aliases):
                matches.append(medication)

        return matches

    def _calculate_data_version(self):
        digest = hashlib.sha256()
        for path in (self.patients_path, self.medications_path):
            with path.open("rb") as stream:
                for chunk in iter(lambda: stream.read(65536), b""):
                    digest.update(chunk)
        return digest.hexdigest()[:12]
